"""Import leads and contacts from a CSV or Excel file (2026-10-06).

The pure part of the Leads / Contacts "Import" steps: reading the file,
guessing which portal box each column belongs to, cleaning the values and
deciding, row by row, what will happen (new / already there / repeat of an
earlier row / skipped). Nothing here touches the database; portal_routes.py
looks up what already exists, shows the Check step and does the saving.

Same rules as the rest of the portal:
  * a phone number is never copied into WhatsApp (or the other way round);
  * a number is kept exactly as written — never given a country code by
    guess; unclear ones are flagged "Check number" by phone_problem();
  * any business's own column names work: every column can be matched by hand.
"""
import csv
import io
import json
import os
import re
import secrets
import time

STAGING_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "import_staging")
MAX_BYTES = 15 * 1024 * 1024
MAX_ROWS = 20000
STAGE_TTL_SECONDS = 24 * 3600

# (key, label) in the order the "Save it as" drop-down shows them.
FIELDS = [
    ("skip",              "Don't import"),
    ("company",           "Company name"),
    ("contact_person",    "Contact Person"),
    ("phone",             "Phone number"),
    ("whatsapp",          "WhatsApp number"),
    ("other_number",      "Other numbers (saved in Notes)"),
    ("email",             "Email"),
    ("website",           "Website"),
    ("street",            "Street address"),
    ("city",              "City / Town"),
    ("state",             "State / County / Province"),
    ("country",           "Country"),
    ("postcode",          "Postcode"),
    ("business_category", "Business category"),
    ("employees",         "Number of employees"),
    ("notes",             "Notes"),
    ("labels",            "Labels"),
    ("deal_value",        "Deal value (leads only)"),
    ("product_interest",  "Product / Service interested in (leads only)"),
]
FIELD_LABELS = dict(FIELDS)
# Several columns may go into these; every other box takes one column.
MULTI_FIELDS = {"skip", "other_number", "notes", "labels"}

_GUESS = {
    "company": ("company", "companyname", "business", "businessname", "organisation", "organization",
                "organisationname", "organizationname", "name", "customer", "customername", "displayname",
                "firm", "account", "accountname"),
    "contact_person": ("contactperson", "contact", "contactname", "companymanager", "manager", "owner",
                       "ceo", "director", "md", "fullname", "person", "keycontact"),
    "phone": ("phone", "phonenumber", "phoneno", "tel", "telephone", "telephonenumber", "number",
              "officephone", "landline", "businessphone"),
    "whatsapp": ("whatsapp", "whatsappnumber", "whatsappno", "whatsappphone", "wa", "wanumber"),
    "other_number": ("mobile", "mobilenumber", "mobileno", "cell", "cellphone", "cellnumber",
                     "additionalnumbers", "additionalnumber", "othernumbers", "othernumber", "otherphone",
                     "phone2", "alternativenumber", "altphone"),
    "email": ("email", "emailaddress", "mail", "emails"),
    "website": ("website", "web", "url", "site", "websiteurl", "webaddress", "homepage"),
    "street": ("street", "streetaddress", "address", "address1", "addressline1", "location"),
    "city": ("city", "town", "citytown", "lga", "area"),
    "state": ("state", "county", "province", "region", "statecountyprovince", "stateprovince"),
    "country": ("country", "countrycode", "nation"),
    "postcode": ("postcode", "postalcode", "zip", "zipcode"),
    "business_category": ("categories", "category", "businesscategory", "industry", "sector",
                          "businesstype", "type", "niche"),
    "employees": ("employees", "numberofemployees", "noofemployees", "employeecount", "staff",
                  "staffsize", "companysize", "size", "headcount", "staffcount"),
    "notes": ("notes", "note", "comment", "comments", "remarks", "description"),
    "labels": ("labels", "label", "tags", "tag"),
    "deal_value": ("dealvalue", "value", "amount", "budget"),
    "product_interest": ("productinterest", "product", "interestedin", "service", "productservice"),
    "skip": ("no", "sn", "sno", "serialnumber", "id", "row", "index"),
}
_HEADER_TO_FIELD = {h: k for k, hs in _GUESS.items() for h in hs}


def norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def guess_mapping(headers: list) -> list:
    """One field key per column. Each one-column box is given to the first
    column that matches it; a later match becomes "Don't import"."""
    out, used = [], set()
    for h in headers:
        k = _HEADER_TO_FIELD.get(norm_header(h), "skip")
        if k not in MULTI_FIELDS and k in used:
            k = "skip"
        used.add(k)
        out.append(k)
    # A file with Mobile but no Phone column: Mobile is the phone number.
    if "phone" not in out:
        for i, h in enumerate(headers):
            if out[i] == "other_number" and norm_header(h).startswith(("mobile", "cell")):
                out[i] = "phone"
                break
    return out


def mapping_problems(mapping: list) -> list:
    """Plain-language problems with a chosen mapping (empty = fine)."""
    probs = []
    for key, label in FIELDS:
        if key in MULTI_FIELDS:
            continue
        n = mapping.count(key)
        if n > 1:
            probs.append(f"\"{label}\" is chosen for {n} columns. Pick it for one column only.")
    if not any(k in mapping for k in ("phone", "whatsapp", "email", "other_number")):
        probs.append("Choose at least one column for Phone number, WhatsApp number or Email. "
                     "A contact needs one of them.")
    if "company" not in mapping and "contact_person" not in mapping:
        probs.append("Choose the column that holds the Company name (or the Contact Person).")
    return probs


# ── Reading the file ──────────────────────────────────────────────────────

def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else repr(v)
    if hasattr(v, "strftime"):
        try:
            return v.strftime("%Y-%m-%d")
        except Exception:
            return str(v)
    return str(v).strip()


def _tidy_table(raw_rows: list):
    """First non-empty row = headers. Drops empty rows and empty trailing
    columns; blank headers become "Column N"."""
    rows = [[_cell_text(c) for c in r] for r in raw_rows]
    rows = [r for r in rows if any(c for c in r)]
    if not rows:
        return [], []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    while width and not any(r[width - 1] for r in rows):
        width -= 1
        rows = [r[:width] for r in rows]
    headers = [h or f"Column {i + 1}" for i, h in enumerate(rows[0])]
    return headers, rows[1:]


def read_file(file_name: str, data: bytes) -> dict:
    """{"file_name", "sheets": [{"name","headers","rows","looks_like_list"}]}.
    Raises ValueError with a message for the person when the file can't be used."""
    name = (file_name or "").strip()
    low = name.lower()
    if not data:
        raise ValueError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ValueError("The file is bigger than 15 MB. Split it into smaller files and import them one by one.")
    sheets = []
    if low.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise ValueError("This Excel file couldn't be opened. Open it in Excel and save it again as .xlsx, "
                             "or save it as CSV.")
        for ws in wb.worksheets:
            raw = []
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                if i > MAX_ROWS + 1:
                    break
                raw.append(list(r))
            headers, rows = _tidy_table(raw)
            sheets.append({"name": ws.title, "headers": headers, "rows": rows})
        wb.close()
    elif low.endswith(".xls"):
        raise ValueError("Old Excel files (.xls) can't be read. In Excel choose File › Save As › "
                         "Excel Workbook (.xlsx) or CSV, then import that file.")
    elif low.endswith((".csv", ".txt")):
        text = None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:5000]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            delim = dialect.delimiter
        except Exception:
            delim = ","
        raw = list(csv.reader(io.StringIO(text), delimiter=delim))[: MAX_ROWS + 2]
        headers, rows = _tidy_table(raw)
        sheets.append({"name": "CSV", "headers": headers, "rows": rows})
    else:
        raise ValueError("Choose a CSV file or an Excel file (.xlsx).")
    for s in sheets:
        guessed = guess_mapping(s["headers"])
        s["looks_like_list"] = bool(s["rows"]) and sum(1 for k in guessed if k != "skip") >= 2
        s["row_count"] = len(s["rows"])
    if not any(s["rows"] for s in sheets):
        raise ValueError("No rows found. The first row must hold the column names and the rows below it the data.")
    if sum(s["row_count"] for s in sheets) > MAX_ROWS:
        raise ValueError(f"The file has more than {MAX_ROWS:,} rows. Split it into smaller files.")
    return {"file_name": name, "sheets": sheets}


# ── Staging between the steps ─────────────────────────────────────────────

def _stage_path(token: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9_-]{20,64}", token or ""):
        raise ValueError("bad token")
    return os.path.join(STAGING_DIR, token + ".json")


def stage_save(data: dict, token: str = None) -> str:
    os.makedirs(STAGING_DIR, exist_ok=True)
    now = time.time()
    for fn in os.listdir(STAGING_DIR):
        p = os.path.join(STAGING_DIR, fn)
        try:
            if now - os.path.getmtime(p) > STAGE_TTL_SECONDS:
                os.remove(p)
        except OSError:
            pass
    token = token or secrets.token_urlsafe(24)
    tmp = _stage_path(token) + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f)
    os.replace(tmp, _stage_path(token))
    return token


def stage_load(token: str):
    try:
        with open(_stage_path(token), encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return None


def stage_delete(token: str):
    try:
        os.remove(_stage_path(token))
    except (OSError, ValueError):
        pass


def combined_table(staged: dict):
    """Headers and rows of the chosen sheets put together. Sheets with the
    same column names line up; a column only one sheet has is added on.
    Each row carries where it came from: (sheet name, row number in file)."""
    chosen = [s for s in staged["sheets"] if s["name"] in (staged.get("chosen_sheets") or [])]
    headers = []
    for s in chosen:
        for h in s["headers"]:
            if h not in headers:
                headers.append(h)
    rows, origin = [], []
    for s in chosen:
        pos = [headers.index(h) for h in s["headers"]]
        for n, r in enumerate(s["rows"]):
            out = [""] * len(headers)
            for i, v in enumerate(r):
                if i < len(pos):
                    out[pos[i]] = v
            rows.append(out)
            origin.append((s["name"], n + 2))
    return headers, rows, origin


# ── Cleaning values ───────────────────────────────────────────────────────

_NUM_SPLIT = re.compile(r"[,;/|]|\s{2,}|\bor\b|\band\b", re.I)


def clean_numbers(value) -> list:
    """Every phone number in a cell, as written: '+' and digits when it was
    written with + or 00, else just the digits. Words like 'no whatsapp' or
    'no cell phone', and anything under 7 digits, are dropped."""
    out = []
    for part in _NUM_SPLIT.split(str(value or "")):
        s = part.strip()
        digits = re.sub(r"\D", "", s)
        if len(digits) < 7:
            continue
        if s.startswith("+"):
            n = "+" + digits
        elif digits.startswith("00"):
            n = "+" + digits[2:]
        else:
            n = digits
        if n not in out:
            out.append(n)
    return out


def number_key(n: str) -> str:
    """What two numbers are compared on: digits only."""
    d = re.sub(r"\D", "", n or "")
    return d[2:] if d.startswith("00") else d


_EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")


def clean_emails(value) -> list:
    return list(dict.fromkeys(m.group(0).lower() for m in _EMAIL_RE.finditer(str(value or ""))))


def clean_text(value, max_len: int):
    v = " ".join(str(value or "").split())
    return v[:max_len] or None


def clean_money(value):
    s = re.sub(r"[^\d.]", "", str(value or ""))
    try:
        return round(float(s), 2) if s else None
    except ValueError:
        return None


def build_records(headers, rows, origin, mapping, default_country, match_country, match_state):
    """One record per row with the chosen columns cleaned."""
    recs = []
    cols = {}
    for i, k in enumerate(mapping):
        cols.setdefault(k, []).append(i)

    def first(key):
        for i in cols.get(key, []):
            return row[i]
        return ""

    for idx, row in enumerate(rows):
        phones = clean_numbers(first("phone"))
        was = clean_numbers(first("whatsapp"))
        others = []
        for i in cols.get("other_number", []):
            others += clean_numbers(row[i])
        emails = []
        for i in cols.get("email", []):
            emails += clean_emails(row[i])
        notes = [row[i].strip() for i in cols.get("notes", []) if row[i].strip()]
        labels = []
        for i in cols.get("labels", []):
            labels += [x.strip()[:50] for x in re.split(r"[,;]", row[i]) if x.strip()]
        country_txt = first("country")
        iso = match_country(country_txt) if country_txt else None
        iso = iso or default_country or None
        state = match_state(iso, first("state")) if (iso and first("state")) else None
        rec = {
            "row": idx,
            "origin": origin[idx],
            "company": clean_text(first("company"), 255),
            "contact_person": clean_text(first("contact_person"), 200),
            "phone": phones[0] if phones else None,
            "whatsapp": was[0] if was else None,
            "email": emails[0] if emails else None,
            "other_emails": emails[1:],
            "website": clean_text(first("website"), 255),
            "street": clean_text(first("street"), 255),
            "city": clean_text(first("city"), 120),
            "state": state,
            "country": iso,
            "postcode": clean_text(first("postcode"), 20),
            "business_category": clean_text(first("business_category"), 120),
            "employees": clean_text(first("employees"), 40),
            "deal_value": clean_money(first("deal_value")),
            "product_interest": clean_text(first("product_interest"), 255),
            "notes": notes,
            "labels": labels,
        }
        # Extra numbers: the rest of the phone/WhatsApp cells plus the other
        # number columns, minus the ones already saved as Phone / WhatsApp.
        kept = {number_key(rec["phone"]), number_key(rec["whatsapp"])}
        extra = []
        for n in phones[1:] + was[1:] + others:
            if number_key(n) not in kept:
                kept.add(number_key(n))
                extra.append(n)
        rec["other_numbers"] = extra
        if not rec["company"]:
            rec["company"] = rec["contact_person"] or rec["email"] or rec["phone"] or rec["whatsapp"]
        if rec["city"] and rec["city"].lower() in ("no city", "n/a", "na", "none", "-"):
            rec["city"] = None
        if rec["employees"] and rec["employees"].lower() in ("n/a", "na", "-", "none"):
            rec["employees"] = None
        recs.append(rec)
    return recs


def record_keys(rec) -> set:
    keys = set()
    for n in (rec.get("phone"), rec.get("whatsapp")):
        if n:
            keys.add("n:" + number_key(n))
    if rec.get("email"):
        keys.add("e:" + rec["email"].lower())
    return keys


_FILL_FIELDS = ("company", "contact_person", "phone", "whatsapp", "email", "website", "street", "city",
                "state", "country", "postcode", "business_category", "employees", "deal_value",
                "product_interest")


def _join_into(first, rec):
    """A later row about the same business fills the first row's empty boxes;
    its other numbers / notes / labels are added."""
    for f in _FILL_FIELDS:
        if not first.get(f) and rec.get(f):
            first[f] = rec[f]
    kept = {number_key(first.get("phone")), number_key(first.get("whatsapp"))} | \
           {number_key(n) for n in first["other_numbers"]}
    for n in [rec.get("phone"), rec.get("whatsapp")] + rec["other_numbers"]:
        if n and number_key(n) not in kept:
            kept.add(number_key(n))
            first["other_numbers"].append(n)
    for e in ([rec["email"]] if rec.get("email") else []) + rec["other_emails"]:
        if e != first.get("email") and e not in first["other_emails"]:
            first["other_emails"].append(e)
    for n in rec["notes"]:
        if n not in first["notes"]:
            first["notes"].append(n)
    for lb in rec["labels"]:
        if lb not in first["labels"]:
            first["labels"].append(lb)
    first.setdefault("joined_rows", []).append(rec["origin"])


def plan(records, existing: dict, phone_problem):
    """Decide what happens to every record.
    existing: {key: id} for what the business already has (same keys as
    record_keys). Returns {"items": [...], "skipped": [...], counts…}.
    Each item gets status 'new' or 'existing' (+ existing_id)."""
    items, skipped, owner = [], [], {}
    joined = 0
    examples = 0
    for rec in records:
        if is_template_example(rec):
            rec["skip_reason"] = "Example row from the template"
            skipped.append(rec)
            examples += 1
            continue
        keys = record_keys(rec)
        if not keys:
            rec["skip_reason"] = "No phone number, WhatsApp number or email"
            skipped.append(rec)
            continue
        hit = next((owner[k] for k in keys if k in owner), None)
        if hit is not None:
            _join_into(hit, rec)
            for k in record_keys(hit) | keys:
                owner.setdefault(k, hit)
            joined += 1
            continue
        for k in keys:
            owner[k] = rec
        items.append(rec)
    for rec in items:
        ex = next((existing[k] for k in sorted(record_keys(rec)) if k in existing), None)
        rec["status"] = "existing" if ex else "new"
        rec["existing_id"] = ex
        rec["check"] = [f for f in ("phone", "whatsapp") if rec.get(f) and phone_problem(rec[f])]
    return {
        "items": items,
        "skipped": skipped,
        "new": sum(1 for r in items if r["status"] == "new"),
        "existing": sum(1 for r in items if r["status"] == "existing"),
        "joined": joined,
        "examples": examples,
        "to_check": sum(1 for r in items if r["check"]),
        "with_whatsapp": sum(1 for r in items if r.get("whatsapp")),
    }


def notes_text(rec) -> str:
    parts = list(rec["notes"])
    if rec["other_numbers"]:
        parts.append("Other numbers: " + ", ".join(rec["other_numbers"]))
    return "\n".join(parts)


def skipped_csv(headers, rows, skipped) -> str:
    """The skipped rows exactly as they were in the file, plus a Reason
    column, so they can be fixed and imported again."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(list(headers) + ["Reason"])
    for rec in skipped:
        w.writerow(list(rows[rec["row"]]) + [rec.get("skip_reason", "")])
    return buf.getvalue()


def check_csv(items, phone_problem) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Company", "Phone number", "WhatsApp number", "Problem", "Sheet", "Row"])
    for rec in items:
        if not rec["check"]:
            continue
        probs = [f"{'Phone' if f == 'phone' else 'WhatsApp'}: {phone_problem(rec[f])}" for f in rec["check"]]
        w.writerow([rec["company"] or "", rec.get("phone") or "", rec.get("whatsapp") or "",
                    "; ".join(probs), rec["origin"][0], rec["origin"][1]])
    return buf.getvalue()


# ── Template to download (2026-10-06) ─────────────────────────────────────
# Column names match _GUESS, so a filled-in template is matched by itself in
# step 2. The two example rows are skipped by the import (user 2026-10-06),
# in case someone forgets to delete them.

TEMPLATE_COLUMNS = ["Company name", "Contact Person", "Phone number", "WhatsApp number", "Other numbers",
                    "Email", "Website", "Street", "City", "State", "Country", "Business category",
                    "Number of employees", "Labels", "Notes"]
TEMPLATE_LEAD_COLUMNS = ["Deal value", "Product / Service"]
TEMPLATE_REQUIRED = {"Phone number", "WhatsApp number", "Email"}
TEMPLATE_EXAMPLES = [
    ["Example Foods Ltd", "Ada Obi", "+234 803 123 4567", "+234 803 123 4567", "+234 701 000 0000",
     "ada@examplefoods.ng", "examplefoods.ng", "12 Allen Avenue", "Ikeja", "Lagos", "Nigeria",
     "Food & Drink", "6-10", "Trade fair", "Met at the expo", "450000", "Catering supplies"],
    ["Example Legal Partners", "", "+44 20 7946 0000", "", "", "info@examplelegal.co.uk", "", "",
     "London", "", "United Kingdom", "Legal", "11-15", "", "", "", ""],
]
_EXAMPLE_COMPANIES = {r[0].lower() for r in TEMPLATE_EXAMPLES}
_EXAMPLE_EMAILS = {r[5].lower() for r in TEMPLATE_EXAMPLES if r[5]}
_EXAMPLE_NUMBERS = {number_key(r[i]) for r in TEMPLATE_EXAMPLES for i in (2, 3, 4) if r[i]}


def is_template_example(rec) -> bool:
    """A row still holding one of the template's example rows: the example
    company name together with its example email or number."""
    if (rec.get("company") or "").strip().lower() not in _EXAMPLE_COMPANIES:
        return False
    if rec.get("email") and rec["email"].lower() in _EXAMPLE_EMAILS:
        return True
    return any(n and number_key(n) in _EXAMPLE_NUMBERS for n in (rec.get("phone"), rec.get("whatsapp")))


def template_columns(target: str) -> list:
    return TEMPLATE_COLUMNS + (TEMPLATE_LEAD_COLUMNS if target == "leads" else [])


def _template_rows(target: str) -> list:
    n = len(template_columns(target))
    return [r[:n] for r in TEMPLATE_EXAMPLES]


def template_howto(target: str) -> list:
    """(is_heading, text) lines for the "How to fill it in" sheet."""
    lines = [
        (True, "Your list: how to fill it in"),
        (False, "Keep row 1 as it is. It holds the column names. Type your businesses or people from row 2, one per row."),
        (False, "Delete the two grey example rows, or type over them. If you leave them in, the import skips them."),
        (True, "Every row needs one of these"),
        (False, "Phone number, WhatsApp number or Email. Rows with none of the three are skipped. "
                "You can download them after the import to fix them."),
        (True, "Phone and WhatsApp numbers"),
        (False, "Write the country code, starting with +. For example +234 803 123 4567 or +44 20 7946 0000."),
        (False, "Numbers without a country code are still imported but marked \"Check number\". "
                "The portal never guesses a country code."),
        (False, "Only put a number under WhatsApp number if that number is on WhatsApp. "
                "The phone number is never copied into WhatsApp."),
        (False, "Extra numbers go in Other numbers, separated by commas. They are saved in Notes."),
        (True, "Everything else"),
        (False, "Leave any column empty if you don't have it. You can also delete columns you don't need."),
        (False, "Business category and Number of employees: write them the way your business uses them, "
                "for example \"Legal\" or \"6-10\"."),
        (False, "Labels: several in one cell, separated by commas, for example \"Trade fair, VIP\"."),
        (False, "Country: the full name (Nigeria) or the 2-letter code (NG)."),
    ]
    if target == "leads":
        lines.append((False, "Deal value: a number only, for example 450000. "
                             "Product / Service: what they are interested in."))
    return lines


def template_csv(target: str) -> str:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(template_columns(target))
    for r in _template_rows(target):
        w.writerow(r)
    return buf.getvalue()


def template_xlsx(target: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Your list"
    cols = template_columns(target)
    ws.append(cols)
    for r in _template_rows(target):
        ws.append(r)
    head = PatternFill("solid", fgColor="E9F3EC")
    req = PatternFill("solid", fgColor="CFE8D6")
    for i, name in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="111E2D")
        c.fill = req if name in TEMPLATE_REQUIRED else head
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(32, len(name) + 6))
    grey = Font(italic=True, color="6B7280")
    for row in ws.iter_rows(min_row=2, max_row=3):
        for c in row:
            c.font = grey
            c.number_format = "@"
    # Text format for the number columns, so Excel never turns +234… into 2.34E+12.
    for i, name in enumerate(cols, 1):
        if name in ("Phone number", "WhatsApp number", "Other numbers"):
            for r in range(2, 2001):
                ws.cell(row=r, column=i).number_format = "@"
    ws.freeze_panes = "A2"
    hs = wb.create_sheet("How to fill it in")
    hs.column_dimensions["A"].width = 110
    for is_head, text in template_howto(target):
        hs.append([text])
        c = hs.cell(row=hs.max_row, column=1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if is_head:
            c.font = Font(bold=True, color="111E2D")
            c.fill = head
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
